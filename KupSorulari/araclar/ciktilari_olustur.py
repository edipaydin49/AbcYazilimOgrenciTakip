"""soru_havuzu.json'dan KupSoruUretici.html ve KupSoruHavuzu.xlsx dosyalarını üretir."""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

BURA = Path(__file__).parent
KOK = BURA.parent
havuz = json.loads((BURA / "soru_havuzu.json").read_text(encoding="utf-8"))

# ---------- HTML ----------
sablon = (BURA / "sablon.html").read_text(encoding="utf-8")
govde = sablon.replace("/*HAVUZ*/", json.dumps(havuz, ensure_ascii=False))
(KOK / "KupSoruUretici.html").write_text('<!doctype html>\n<html lang="tr">\n' + govde + "</html>\n", encoding="utf-8")
if len(sys.argv) > 1:  # yayın kopyası (iskeletsiz)
    Path(sys.argv[1]).write_text(govde.replace('<meta charset="utf-8">\n', "").replace(
        '<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">\n', ""), encoding="utf-8")

# ---------- Excel ----------
SATIR = 1000          # havuz sayfasında formül kapsayan satır (öğretmen yeni soru ekleyebilsin)
MAKS = 100            # testte en fazla soru
F = "Arial"
HARF = "ABCD"
ince = Side(style="thin", color="999999")
kutu = Border(left=ince, right=ince, top=ince, bottom=ince)
sari = PatternFill("solid", fgColor="FFFF00")
baslik_dolgu = PatternFill("solid", fgColor="1F4E78")
ortala = Alignment(horizontal="center", vertical="center", wrap_text=True)
sar = Alignment(vertical="center", wrap_text=True)

# Şekle bakmadan çözülemeyen sorular Excel'e alınmaz (bunlar yalnızca web sayfasında).
excel_havuz = [s for s in havuz if not (s.get("sekil", {}).get("tip") == "yapi" and s["soru"].startswith("Şekildeki"))]
konular = list(dict.fromkeys(s["konu"] for s in excel_havuz))

wb = Workbook()
ts = wb.active
ts.title = "Test"
hv = wb.create_sheet("Soru Havuzu")
ca = wb.create_sheet("Cevap Anahtarı")


def baslik(ws, satir, basliklar):
    for j, h in enumerate(basliklar, start=1):
        c = ws.cell(satir, j, h)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = baslik_dolgu
        c.alignment = ortala
        c.border = kutu


# --- Soru Havuzu ---
baslik(hv, 1, ["No", "Konu", "Zorluk", "Soru", "A", "B", "C", "D", "Doğru Şık", "Uygun", "Rastgele", "Sıra"])
for i, s in enumerate(excel_havuz, start=2):
    deger = [i - 1, s["konu"], s["zorluk"], s["soru"], *s["secenekler"], HARF[s["dogru"]]]
    for j, v in enumerate(deger, start=1):
        c = hv.cell(i, j, v)
        c.font = Font(name=F)
        c.alignment = sar if j == 4 else ortala
for r in range(2, SATIR + 2):
    hv.cell(r, 10, f'=IF($D{r}="",0,IF(AND(OR(Test!$B$4="Hepsi",$B{r}=Test!$B$4),OR(Test!$B$5="Hepsi",$C{r}=Test!$B$5)),1,0))')
    hv.cell(r, 11, f"=IF($J{r}=1,RAND(),-1)")
    hv.cell(r, 12, f'=IF($J{r}=1,RANK($K{r},$K$2:$K${SATIR + 1}),"")')
    for j in (10, 11, 12):
        hv.cell(r, j).font = Font(name=F, color="808080")
for col, w in zip("ABCDEFGHIJKL", (6, 22, 9, 70, 16, 16, 16, 16, 10, 8, 10, 7)):
    hv.column_dimensions[col].width = w
hv.freeze_panes = "A2"
hv["D1"].comment = Comment("Yeni soru eklemek için son sorunun altındaki boş satıra yazın: Konu, Zorluk "
                           "(Kolay/Orta/Zor), Soru, A-D seçenekleri ve Doğru Şık. Gri sütunlara dokunmayın.", "Not")

# --- Test ---
ts["A1"] = "6. SINIF KÜP SORULARI · RASTGELE TEST"
ts["A1"].font = Font(name=F, size=16, bold=True, color="1F4E78")
ts.merge_cells("A1:H1")
ayarlar = [("Soru sayısı", 10, f"1 ile {MAKS} arası"),
           ("Konu", "Hepsi", "Listeden seçin"),
           ("Zorluk", "Hepsi", "Hepsi / Kolay / Orta / Zor")]
for i, (etiket, deger, not_) in enumerate(ayarlar, start=3):
    ts.cell(i, 1, etiket).font = Font(name=F, bold=True)
    c = ts.cell(i, 2, deger)
    c.fill = sari
    c.font = Font(name=F, bold=True, color="0000FF")
    c.alignment = ortala
    c.border = kutu
    ts.cell(i, 3, not_).font = Font(name=F, italic=True, color="666666")
ts["A6"] = "Uygun soru (havuzda)"
ts["A6"].font = Font(name=F, bold=True)
ts["B6"] = f"=SUM('Soru Havuzu'!$J$2:$J${SATIR + 1})"
ts["B6"].font = Font(name=F, bold=True)
ts["B6"].alignment = ortala
ts["C6"] = '=IF(B6<B3,"Uyarı: havuzda yeterli soru yok, "&B6&" soru listelendi.","")'
ts["C6"].font = Font(name=F, bold=True, color="C00000")

dv_adet = DataValidation(type="whole", operator="between", formula1="1", formula2=str(MAKS), showErrorMessage=True,
                         errorTitle="Hatalı değer", error=f"1 ile {MAKS} arasında bir sayı girin.")
dv_konu = DataValidation(type="list", formula1='"' + ",".join(["Hepsi"] + konular) + '"')
dv_zor = DataValidation(type="list", formula1='"Hepsi,Kolay,Orta,Zor"')
for d in (dv_adet, dv_konu, dv_zor):
    ts.add_data_validation(d)
dv_adet.add("B3"); dv_konu.add("B4"); dv_zor.add("B5")

ts["A8"] = "Sarı hücreleri değiştirin. Yeni sorular için F9'a basın. Listeyi sabitlemek için: Kopyala → Değerleri Yapıştır."
ts["A8"].font = Font(name=F, italic=True, color="C00000")
ts.merge_cells("A8:H8")
ts["A9"] = "Ad Soyad: .................................     Sınıf: ...........     Tarih: ....../....../......"
ts["A9"].font = Font(name=F)
ts.merge_cells("A9:H9")
baslik(ts, 11, ["No", "Soru", "A)", "B)", "C)", "D)", "Konu", "Zorluk"])
ts.cell(11, 10, "Havuz satırı").font = Font(name=F, color="808080")

ilk = 12
aralik = f"'Soru Havuzu'!$L$2:$L${SATIR + 1}"
for k in range(1, MAKS + 1):
    r = ilk + k - 1
    goster = f"{k}<=MIN($B$3,$B$6)"
    ts.cell(r, 10, f'=IF({goster},MATCH({k},{aralik},0),"")').font = Font(name=F, color="808080")
    ts.cell(r, 1, f'=IF($J{r}="","",{k})')
    for j, kaynak in ((2, "D"), (3, "E"), (4, "F"), (5, "G"), (6, "H"), (7, "B"), (8, "C")):
        ts.cell(r, j, f"=IF($J{r}=\"\",\"\",INDEX('Soru Havuzu'!${kaynak}$2:${kaynak}${SATIR + 1},$J{r}))")
    for j in range(1, 9):
        c = ts.cell(r, j)
        c.font = Font(name=F, size=11, bold=(j == 1))
        c.alignment = sar if j == 2 else ortala
son = ilk + MAKS - 1
ts.conditional_formatting.add(f"A{ilk}:H{son}", FormulaRule(formula=[f'$A{ilk}<>""'], border=kutu))
for col, w in zip("ABCDEFGHIJ", (5, 60, 13, 13, 13, 13, 18, 8, 3, 11)):
    ts.column_dimensions[col].width = w
ts.column_dimensions["I"].hidden = True
ts.freeze_panes = f"A{ilk}"
ts.print_title_rows = "11:11"
ts.page_setup.orientation = "landscape"
ts.page_setup.fitToWidth = 1
ts.page_setup.fitToHeight = 0
ts.sheet_properties.pageSetUpPr.fitToPage = True

# --- Cevap Anahtarı ---
ca["A1"] = "CEVAP ANAHTARI"
ca["A1"].font = Font(name=F, size=14, bold=True, color="1F4E78")
baslik(ca, 3, ["No", "Doğru Şık", "Cevap"])
for k in range(1, MAKS + 1):
    r = 3 + k
    t = ilk + k - 1
    ca.cell(r, 1, f"=Test!A{t}")
    ca.cell(r, 2, f"=IF(Test!$J{t}=\"\",\"\",INDEX('Soru Havuzu'!$I$2:$I${SATIR + 1},Test!$J{t}))")
    ca.cell(r, 3, f'=IF(B{r}="","",INDEX(Test!$C{t}:$F{t},MATCH(B{r},{{"A","B","C","D"}},0)))')
    for j in (1, 2, 3):
        ca.cell(r, j).font = Font(name=F, bold=(j == 2))
        ca.cell(r, j).alignment = ortala
ca.column_dimensions["A"].width = 6
ca.column_dimensions["B"].width = 11
ca.column_dimensions["C"].width = 22

wb.save(KOK / "KupSoruHavuzu.xlsx")
print("HTML:", len(havuz), "soru · Excel:", len(excel_havuz), "soru")
